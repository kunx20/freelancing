import csv
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

total_revenue = 0
total_quantity = 0
total_records = 0

product_summary = defaultdict(
    lambda: {"quantity": 0, "revenue": 0}
)

sales_details = []

# Read validated sales data
with open("validated_sales.csv", "r", newline="") as file:
    reader = csv.DictReader(file)

    for row in reader:
        product = row["Product"]
        quantity = int(row["Quantity"])
        price = float(row["Price"])

        revenue = quantity * price

        total_revenue += revenue
        total_quantity += quantity
        total_records += 1

        product_summary[product]["quantity"] += quantity
        product_summary[product]["revenue"] += revenue

        sales_details.append({
            "Product": product,
            "Quantity": quantity,
            "Price": price,
            "Revenue": revenue
        })

# Create Excel workbook
workbook = Workbook()

# Worksheet 1: Summary
summary = workbook.active
summary.title = "Summary"

summary["A1"] = "Sales Report"
summary["A1"].font = Font(bold=True, size=16)

summary["A3"] = "Total Revenue"
summary["B3"] = total_revenue

summary["A4"] = "Total Quantity Sold"
summary["B4"] = total_quantity

summary["A5"] = "Valid Sales Records"
summary["B5"] = total_records

# Product-wise summary
summary.append([])
summary.append(["Product", "Total Quantity", "Total Revenue"])

for product, details in product_summary.items():
    summary.append([
        product,
        details["quantity"],
        details["revenue"]
    ])

# Worksheet 2: Sales Details
details_sheet = workbook.create_sheet("Sales Details")

details_sheet.append([
    "Product",
    "Quantity",
    "Unit Price",
    "Revenue"
])

for record in sales_details:
    details_sheet.append([
        record["Product"],
        record["Quantity"],
        record["Price"],
        record["Revenue"]
    ])

# Format worksheet headers
header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

for sheet, header_row in [
    (summary, 7),
    (details_sheet, 1)
]:
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

# Format currency values in the Summary worksheet
summary["B3"].number_format = '#,##0.00'

for row in range(8, summary.max_row + 1):
    summary.cell(row=row, column=3).number_format = '#,##0.00'

# Format currency values in Sales Details
for row in range(2, details_sheet.max_row + 1):
    details_sheet.cell(row=row, column=3).number_format = '#,##0.00'
    details_sheet.cell(row=row, column=4).number_format = '#,##0.00'

# Adjust column widths automatically
for sheet in workbook.worksheets:
    for column in sheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )

        sheet.column_dimensions[
            column[0].column_letter
        ].width = max_length + 3

    sheet.freeze_panes = "A2"

# Save Excel report
workbook.save("sales_report.xlsx") 

# Professional formatting
header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

thin_border = Border(
    bottom=Side(style="thin", color="D9E2F3")
)

# Format the Summary worksheet
for cell in summary[7]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

# Format the Sales Details worksheet
for cell in details_sheet[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

# Add borders, filters and suitable column widths
for sheet in workbook.worksheets:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = thin_border

    for column in sheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )

        sheet.column_dimensions[
            column[0].column_letter
        ].width = min(max_length + 3, 35)

# Save the Excel report
workbook.save("sales_report.xlsx")


print("Excel report generated successfully!")
print("Total revenue:", total_revenue)
print("Total quantity sold:", total_quantity)
print("File created: sales_report.xlsx")